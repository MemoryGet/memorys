"""Real converter probe in an explicitly provisioned venv; never production acceptance."""
import argparse
import base64
import hashlib
import io
import json
from pathlib import Path
import subprocess
import time
import tempfile
import sys
import zipfile

ROOT = Path(__file__).resolve().parents[2]


def fixtures():
    output = io.BytesIO()
    with zipfile.ZipFile(output, "w") as package:
        package.writestr("[Content_Types].xml", '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        package.writestr("_rels/.rels", '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        package.writestr("word/document.xml", '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>中文设备维护知识</w:t></w:r></w:p></w:body></w:document>')
    docx = output.getvalue()
    from openpyxl import Workbook
    workbook = Workbook()
    workbook.active.append(["设备", "温度"])
    workbook.active.append(["泵站", 90])
    output = io.BytesIO()
    workbook.save(output)
    xlsx = output.getvalue()
    from pptx import Presentation
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "中文行业知识"
    slide.placeholders[1].text = "设备检查流程"
    output = io.BytesIO()
    presentation.save(output)
    pptx = output.getvalue()
    content = b"BT /F1 12 Tf 72 720 Td (Hello PDF knowledge) Tj ET"
    objects = [b"<< /Type /Catalog /Pages 2 0 R >>", b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>", b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>", b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>", b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream"]
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, 1):
        offsets.append(len(pdf))
        pdf += str(number).encode() + b" 0 obj\n" + obj + b"\nendobj\n"
    start = len(pdf)
    pdf += b"xref\n0 6\n0000000000 65535 f \n"
    pdf += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    pdf += f"trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n{start}\n%%EOF\n".encode()
    return {"docx": (docx, "中文设备维护知识"), "xlsx": (xlsx, "泵站"), "pptx": (pptx, "设备检查流程"), "pdf": (bytes(pdf), "Hello PDF knowledge")}


def run_probe():
    from check_architecture import source_inventory
    before = source_inventory()
    worker = ROOT / "tools/parsers/markitdown_worker.py"
    records = []
    parent = ROOT / ".local/architecture-check"
    parent.mkdir(parents=True, exist_ok=True)
    temporary = tempfile.TemporaryDirectory(prefix="office-probe-", dir=parent)
    try:
        return run_in_directory(Path(temporary.name), worker, records, before)
    finally:
        temporary.cleanup()


def run_in_directory(folder, worker, records, before):
    from check_architecture import source_inventory
    executable = folder / "office-parser-probe.exe"
    subprocess.run([str(ROOT / ".tools/go/bin/go.exe"), "build", "-o", str(executable), "./cmd/architecture-office-parser"], cwd=ROOT, check=True, timeout=90)
    executable_sha256 = hashlib.sha256(executable.read_bytes()).hexdigest()
    for format_name, (raw, expected) in fixtures().items():
        request = {"format": format_name, "data_base64": base64.b64encode(raw).decode(), "sha256": hashlib.sha256(raw).hexdigest()}
        start = time.monotonic()
        result = subprocess.run([sys.executable, "-I", str(worker)], input=json.dumps(request), text=True, capture_output=True, encoding="utf-8", timeout=45)
        if result.returncode:
            raise RuntimeError(f"{format_name}: converter rejected fixture: {result.stderr[:500]}")
        parsed = json.loads(result.stdout)
        if parsed["version"] != "markitdown/0.1.3" or expected not in parsed["markdown"]:
            raise RuntimeError(f"{format_name}: expected content missing")
        source = folder / ("source." + format_name)
        source.write_bytes(raw)
        go_result = subprocess.run([str(executable), "--python", sys.executable, "--worker", str(worker), "--source", str(source), "--format", format_name, "--root", str(folder / "blobs")], capture_output=True, text=True, encoding="utf-8", timeout=45)
        if go_result.returncode:
            raise RuntimeError(f"{format_name}: Go parser chain failed: {go_result.stderr[:500]}")
        go_evidence = json.loads(go_result.stdout)
        if (go_evidence.get("passed") is not True or go_evidence.get("production_acceptance") is not False
                or go_evidence.get("configured_factory") is not True
                or go_evidence.get("closed_reference_rejected") is not True
                or go_evidence.get("format") != format_name):
            raise RuntimeError(f"{format_name}: Go assembly evidence incomplete")
        records.append({"go_parser_chain": True, "configured_factory": True, "closed_reference_rejected": True, "format": format_name, "source_sha256": request["sha256"], "output_sha256": hashlib.sha256(parsed["markdown"].encode()).hexdigest(), "elapsed_seconds": round(time.monotonic() - start, 3), "contains_expected": True})
    if before != source_inventory():
        raise RuntimeError("probe sources changed during validation")
    import importlib.metadata
    report = {"passed": True, "production_acceptance": False, "scope": "Four synthetic source fixtures; Office Chinese text and ASCII born-digital PDF", "version": importlib.metadata.version("markitdown"), "worker_sha256": hashlib.sha256(worker.read_bytes()).hexdigest(), "records": records, "unverified": ["scanned/Chinese/multicolumn PDF", "precise page/cell provenance and structured tables", "process-tree/resource isolation", "real authorization and lifecycle integration", "full dependency licenses/artifact identities"]}
    report.update(sources=before, sources_unchanged_during_checks=True, probe_executable_sha256=executable_sha256)
    report["dependency_lock_sha256"] = hashlib.sha256((ROOT / "docs/architecture/markitdown-requirements.lock.txt").read_bytes()).hexdigest()
    return report


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--report", type=Path, default=ROOT / ".local/architecture-check/markitdown-runtime.json")
    args = parser.parse_args()
    args.report.parent.mkdir(parents=True, exist_ok=True)
    report = {"passed": False, "production_acceptance": False, "status": "running"}
    args.report.write_text(json.dumps(report) + "\n", encoding="utf-8")
    try:
        report = run_probe()
    except Exception as error:
        report.update(status="failed", error=type(error).__name__)
        args.report.write_text(json.dumps(report) + "\n", encoding="utf-8")
        raise
    report["temporary_resources_removed"] = True
    args.report.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print("MarkItDown real converter fixtures: passed")


if __name__ == "__main__":
    main()
