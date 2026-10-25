"""Create synthetic, non-confidential original files for the isolated CLI test."""
from pathlib import Path
import zipfile
import sys


def create(folder):
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches

    folder.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(folder / 'sample.docx', 'w') as package:
        package.writestr('[Content_Types].xml', '<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"><Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/><Default Extension="xml" ContentType="application/xml"/><Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/></Types>')
        package.writestr('_rels/.rels', '<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"><Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/></Relationships>')
        package.writestr('word/document.xml', '<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"><w:body><w:p><w:r><w:t>localdocxprobe</w:t></w:r></w:p><w:p><w:r><w:t>星帆项目的紧急联络暗号是紫色海豚。遇到数据库故障，先停止写入，再恢复数据库与匹配的本地状态目录。</w:t></w:r></w:p></w:body></w:document>')
    book = Workbook()
    book.active.append(['localxlsxprobe', '联系人', '恢复顺序'])
    book.active.append(['测试记录', '林工', '先数据库，再状态目录'])
    book.save(folder / 'sample.xlsx')
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    frame = slide.shapes.add_textbox(Inches(1), Inches(1), Inches(8), Inches(3)).text_frame
    frame.text = 'localpptxprobe\n月度演练安排在每月第二个周五。'
    deck.save(folder / 'sample.pptx')
    # A real one-page PDF with an embedded text stream and valid xref offsets.
    stream = b'BT /F1 12 Tf 72 720 Td (localpdfprobe recovery contact is Alice.) Tj ET'
    objects = [b'<< /Type /Catalog /Pages 2 0 R >>',
               b'<< /Type /Pages /Kids [3 0 R] /Count 1 >>',
               b'<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>',
               b'<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>',
               b'<< /Length ' + str(len(stream)).encode() + b' >>\nstream\n' + stream + b'\nendstream']
    raw, offsets = b'%PDF-1.4\n', [0]
    for i, obj in enumerate(objects, 1):
        offsets.append(len(raw))
        raw += str(i).encode() + b' 0 obj\n' + obj + b'\nendobj\n'
    xref = len(raw)
    raw += b'xref\n0 6\n0000000000 65535 f \n'
    raw += b''.join(f'{off:010d} 00000 n \n'.encode() for off in offsets[1:])
    raw += b'trailer\n<< /Size 6 /Root 1 0 R >>\nstartxref\n' + str(xref).encode() + b'\n%%EOF\n'
    (folder / 'sample.pdf').write_bytes(raw)


if __name__ == '__main__':
    create(Path(sys.argv[1]))
